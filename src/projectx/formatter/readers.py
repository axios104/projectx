from abc import ABC, abstractmethod
from typing import Iterator, Union, List, Dict, Any, Optional
from pathlib import Path
import json

class BaseReader(ABC):
    """Abstract base reader."""
    @abstractmethod
    def read(self, source: Union[str, Path]) -> List[Dict[str, Any]]: ...
    
    @abstractmethod  
    def read_stream(self, source: Union[str, Path]) -> Iterator[Dict[str, Any]]: ...

class JsonReader(BaseReader):
    """Reads JSON files (array of objects or newline-delimited JSON)."""
    
    def read(self, source: Union[str, Path]) -> List[Dict[str, Any]]:
        return list(self.read_stream(source))

    def read_stream(self, source: Union[str, Path]) -> Iterator[Dict[str, Any]]:
        path = Path(source)
        with open(path, 'r', encoding='utf-8-sig') as f:
            content = f.read().strip()
            
            # Auto-detect array vs JSONL
            if content.startswith('['):
                data = json.loads(content)
                for item in data:
                    yield item
            else:
                f.seek(0)
                for line in f:
                    line = line.strip()
                    if line:
                        yield json.loads(line)

class ExcelReader(BaseReader):
    """Reads Excel (.xlsx) files using openpyxl."""
    
    def __init__(self, sheet_name: Optional[str] = None):
        self.sheet_name = sheet_name

    def read(self, source: Union[str, Path]) -> List[Dict[str, Any]]:
        return list(self.read_stream(source))
        
    def read_stream(self, source: Union[str, Path]) -> Iterator[Dict[str, Any]]:
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl is required to read Excel files. Install it with: pip install openpyxl")
            
        path = Path(source)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        
        if self.sheet_name:
            if self.sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet '{self.sheet_name}' not found.")
            sheet = wb[self.sheet_name]
        else:
            sheet = wb.active
            
        headers = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h) if h is not None else f"col_{j}" for j, h in enumerate(row)]
                continue
            
            record = {}
            for j, val in enumerate(row):
                if j < len(headers):
                    record[headers[j]] = val
            yield record
            
        wb.close()
